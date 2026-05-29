#!/usr/bin/env python3
"""
Manuscript Review Bridge

Purpose:
    Automate the local DOCX -> Google Docs review loop for versioned
    manuscript builds.

Commands:
    init-config: Create a starter config JSON file.
    push:        Upload latest manuscript DOCX as Google Doc, share, archive old.
    restore:     Restore an archived Google Doc back to ACTIVE.
    download:    Export active Google Doc as DOCX and write gdoc_diff.md.
    pull:        Fetch comments and build actionable local review files.
    packet:      Build a combined review packet using comments + local changelog.
    close:       Append response notes for addressed comment IDs.

Examples:
    python scripts/manuscript/manuscript_review_bridge.py init-config
    python scripts/manuscript/manuscript_review_bridge.py push
    python scripts/manuscript/manuscript_review_bridge.py download
    python scripts/manuscript/manuscript_review_bridge.py pull
    python scripts/manuscript/manuscript_review_bridge.py packet
    python scripts/manuscript/manuscript_review_bridge.py close --ids CMT-v034-001 CMT-v034-002
"""

import argparse
import difflib
import html
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

project_root = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(project_root))


##### CONSTANTS #####

MANUSCRIPT_ROOT = project_root / "reports_ARCHIVE" / "manuscript"
MANUSCRIPT_VERSIONS_DIR = MANUSCRIPT_ROOT / "versions"
REVIEW_CONFIG_PATH = MANUSCRIPT_ROOT / "review" / "review_bridge_config.json"
DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


##### DATA STRUCTURES #####

@dataclass
class ReviewConfig:
    """User-facing configuration for Google Docs review automation."""

    credentials_file: str = "secrets/google_credentials.json"
    token_file: str = "secrets/google_token.json"
    drive_parent_folder_id: str = ""
    active_folder_name: str = "ACTIVE"
    archive_folder_name: str = "ARCHIVED"
    reviewer_emails: list[str] = field(default_factory=list)
    share_role: str = "commenter"
    share_with_anyone_with_link: bool = False
    anyone_role: str = "commenter"
    oauth_open_browser: bool = False
    carry_forward_fixed_comments: bool = False
    carry_forward_max_comments_per_doc: int = 30
    auto_refresh_review_docx: bool = True
    auto_compare_active_docx: bool = True
    compare_fail_on_diff: bool = False
    compare_write_report: bool = True
    active_file_id: str = ""


@dataclass
class ReviewDocMeta:
    """Metadata linking a local manuscript version to Google Docs."""

    manuscript_version: int
    local_docx_path: str
    local_version_dir: str
    google_file_id: str
    google_doc_url: str
    uploaded_at_utc: str
    archived_previous_ids: list[str] = field(default_factory=list)
    status: str = "active"


@dataclass
class CommentReply:
    """A reply on a comment thread."""

    reply_id: str
    author: str
    created_time: str
    content: str


@dataclass
class ReviewComment:
    """Normalized comment thread model from Google Drive comments API."""

    comment_id: str
    stable_id: str
    author: str
    created_time: str
    modified_time: str
    resolved: bool
    content: str
    quoted_anchor: str
    section_guess: str
    replies: list[CommentReply] = field(default_factory=list)


##### BASIC HELPERS #####

def _utc_now() -> str:
    """Return an ISO-8601 timestamp in UTC."""
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _extract_version(path: Path) -> int:
    """Extract manuscript version from path stem or parent folder."""
    stem_match = re.search(r"_v(\d+)$", path.stem)
    if stem_match:
        return int(stem_match.group(1))
    dir_match = re.search(r"^v(\d+)$", path.parent.name)
    if dir_match:
        return int(dir_match.group(1))
    return 0


def _collect_existing_manuscripts() -> dict[int, Path]:
    """Find all versioned manuscript DOCX files."""
    found: dict[int, Path] = {}
    for path in sorted(MANUSCRIPT_VERSIONS_DIR.glob("v*/GSM_Manuscript_v*.docx")):
        if path.stem.endswith("_review"):
            continue
        version = _extract_version(path)
        if version > 0:
            found[version] = path
    return found


def _find_manuscript(version: int | None) -> tuple[int, Path]:
    """Return target manuscript version and path."""
    found = _collect_existing_manuscripts()
    if not found:
        raise FileNotFoundError("No manuscript files found in manuscript/versions")
    if version is not None:
        if version not in found:
            raise FileNotFoundError(f"Manuscript version not found: v{version:03d}")
        return version, found[version]
    latest = max(found)
    return latest, found[latest]


def _review_dir_for_version(version: int) -> Path:
    """Return review artifact directory for a manuscript version."""
    review_dir = MANUSCRIPT_VERSIONS_DIR / f"v{version:03d}" / "review"
    review_dir.mkdir(parents=True, exist_ok=True)
    return review_dir


def _safe_read_text(path: Path, default: str = "") -> str:
    """Read text if the file exists, otherwise return default."""
    if not path.exists():
        return default
    return path.read_text(encoding="utf-8")


def _normalize_comment_text(text: str) -> str:
    """Decode HTML entities and collapse whitespace for readability/search."""
    decoded = html.unescape(text or "")
    return re.sub(r"\s+", " ", decoded).strip()


def _derive_search_key(anchor_text: str, max_words: int = 12) -> str:
    """Derive a compact searchable phrase from anchor text.

    We avoid very short/weak fragments and keep a phrase that is likely to
    appear unchanged in the latest manuscript.
    """
    cleaned = _normalize_comment_text(anchor_text)
    if not cleaned:
        return ""
    words = [w for w in re.split(r"\s+", cleaned) if w]
    if not words:
        return ""

    # Prefer a centered chunk; edges are often truncated in API snippets.
    take = min(max_words, len(words))
    start = max(0, (len(words) // 2) - (take // 2))
    chunk = words[start:start + take]

    # If chunk is too short after cleanup, fallback to first words.
    phrase = " ".join(chunk).strip(" .,;:()[]{}")
    if len(phrase) < 20:
        phrase = " ".join(words[:take]).strip(" .,;:()[]{}")
    return phrase


def _extract_drive_id(value: str) -> str:
    """Extract a Drive file id from a URL or raw id string."""
    if not value:
        return ""
    match = re.search(r"/d/([a-zA-Z0-9_-]+)", value)
    if match:
        return match.group(1)
    return value.strip()


##### CONFIG #####

def _config_template() -> ReviewConfig:
    """Return default review bridge configuration."""
    return ReviewConfig(
        credentials_file="secrets/google_credentials.json",
        token_file="secrets/google_token.json",
        drive_parent_folder_id="",
        active_folder_name="ACTIVE",
        archive_folder_name="ARCHIVED",
        reviewer_emails=["instructor1@example.com", "instructor2@example.com"],
        share_role="commenter",
        share_with_anyone_with_link=False,
        anyone_role="commenter",
        oauth_open_browser=False,
        carry_forward_fixed_comments=False,
        carry_forward_max_comments_per_doc=30,
        auto_refresh_review_docx=True,
        auto_compare_active_docx=True,
        compare_fail_on_diff=False,
        compare_write_report=True,
        active_file_id="",
    )


def _load_config() -> ReviewConfig:
    """Load user config from JSON file."""
    if not REVIEW_CONFIG_PATH.exists():
        raise FileNotFoundError(
            "Review config missing. Run init-config first: "
            "python scripts/manuscript/manuscript_review_bridge.py init-config"
        )
    data = json.loads(REVIEW_CONFIG_PATH.read_text(encoding="utf-8"))
    return ReviewConfig(
        credentials_file=data.get("credentials_file", "secrets/google_credentials.json"),
        token_file=data.get("token_file", "secrets/google_token.json"),
        drive_parent_folder_id=data.get("drive_parent_folder_id", ""),
        active_folder_name=data.get("active_folder_name", "ACTIVE"),
        archive_folder_name=data.get("archive_folder_name", "ARCHIVED"),
        reviewer_emails=list(data.get("reviewer_emails", [])),
        share_role=data.get("share_role", "commenter"),
        share_with_anyone_with_link=bool(data.get("share_with_anyone_with_link", False)),
        anyone_role=data.get("anyone_role", "commenter"),
        oauth_open_browser=bool(data.get("oauth_open_browser", False)),
        carry_forward_fixed_comments=bool(data.get("carry_forward_fixed_comments", False)),
        carry_forward_max_comments_per_doc=int(data.get("carry_forward_max_comments_per_doc", 30)),
        auto_refresh_review_docx=bool(data.get("auto_refresh_review_docx", True)),
        auto_compare_active_docx=bool(data.get("auto_compare_active_docx", True)),
        compare_fail_on_diff=bool(data.get("compare_fail_on_diff", False)),
        compare_write_report=bool(data.get("compare_write_report", True)),
        active_file_id=data.get("active_file_id", ""),
    )


def cmd_init_config(_: argparse.Namespace) -> int:
    """Create starter config file if it does not exist."""
    REVIEW_CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
    if REVIEW_CONFIG_PATH.exists():
        print(f"Config already exists: {REVIEW_CONFIG_PATH}")
        return 0
    cfg = _config_template()
    REVIEW_CONFIG_PATH.write_text(
        json.dumps(asdict(cfg), indent=2) + "\n", encoding="utf-8"
    )
    print(f"Created config: {REVIEW_CONFIG_PATH}")
    print("Edit it before running push/pull.")
    return 0


##### GOOGLE CLIENT #####

def _google_modules() -> tuple[Any, Any, Any, Any]:
    """Import Google client libraries lazily with clear install error."""
    try:
        from google.auth.transport.requests import Request
        from google.oauth2.credentials import Credentials
        from google_auth_oauthlib.flow import InstalledAppFlow
        from googleapiclient.discovery import build
    except Exception as exc:
        raise RuntimeError(
            "Missing Google API libraries. Install: "
            "pip install google-api-python-client google-auth-httplib2 "
            "google-auth-oauthlib"
        ) from exc
    return Request, Credentials, InstalledAppFlow, build


def _get_drive_service(config: ReviewConfig):
    """Build authenticated Google Drive API client."""
    Request, Credentials, InstalledAppFlow, build = _google_modules()
    scopes = [
        "https://www.googleapis.com/auth/drive",
        "https://www.googleapis.com/auth/documents.readonly",
    ]
    credentials_path = project_root / config.credentials_file
    token_path = project_root / config.token_file
    token_path.parent.mkdir(parents=True, exist_ok=True)
    creds = None
    if token_path.exists():
        creds = Credentials.from_authorized_user_file(str(token_path), scopes)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            try:
                creds.refresh(Request())
            except Exception as e:
                print(f"Failed to refresh token ({e}). Re-authorizing...")
                creds = None
                
        if not creds or not creds.valid:
            flow = InstalledAppFlow.from_client_secrets_file(str(credentials_path), scopes)
            if config.oauth_open_browser:
                creds = flow.run_local_server(port=0)
            else:
                flow.redirect_uri = "urn:ietf:wg:oauth:2.0:oob"
                auth_url, _ = flow.authorization_url(prompt='consent')
                print(f"\nPlease visit this URL to authorize the application:\n\n{auth_url}\n")
                code = input("Enter the authorization code: ")
                flow.fetch_token(code=code)
                creds = flow.credentials
                
        token_path.write_text(creds.to_json(), encoding="utf-8")
    return build("drive", "v3", credentials=creds)


def _export_google_doc_as_docx(drive, file_id: str, output_path: Path) -> None:
    """Export a Google Doc as a DOCX file."""
    from io import BytesIO

    try:
        from googleapiclient.http import MediaIoBaseDownload
    except Exception as exc:
        raise RuntimeError(
            "Missing Google API libraries. Install: "
            "pip install google-api-python-client google-auth-httplib2 "
            "google-auth-oauthlib"
        ) from exc

    request = drive.files().export_media(
        fileId=file_id,
        mimeType=DOCX_MIME_TYPE,
    )
    buffer = BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(buffer.getvalue())


def _ensure_folder_id(drive, folder_name: str, parent_id: str) -> str:
    """Create or reuse a Drive folder under parent."""
    safe_name = folder_name.replace("'", "\\'")
    query = (
        "mimeType='application/vnd.google-apps.folder' and "
        f"name='{safe_name}' and trashed=false and '{parent_id}' in parents"
    )
    result = drive.files().list(q=query, fields="files(id,name)", pageSize=10).execute()
    files = result.get("files", [])
    if files:
        return files[0]["id"]
    body = {
        "name": folder_name,
        "mimeType": "application/vnd.google-apps.folder",
        "parents": [parent_id],
    }
    created = drive.files().create(body=body, fields="id").execute()
    return created["id"]


def _upload_docx_as_google_doc(drive, file_path: Path, folder_id: str) -> tuple[str, str]:
    """Upload DOCX and convert it to Google Docs format."""
    _, _, _, _ = _google_modules()
    from googleapiclient.http import MediaFileUpload

    body = {
        "name": file_path.stem,
        "parents": [folder_id],
        "mimeType": "application/vnd.google-apps.document",
    }
    media = MediaFileUpload(
        str(file_path),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        resumable=False,
    )
    created = drive.files().create(body=body, media_body=media, fields="id,webViewLink").execute()
    return created["id"], created.get("webViewLink", "")


def _share_with_reviewers(drive, file_id: str, emails: list[str], role: str) -> None:
    """Share a Google Doc with reviewer emails."""
    try:
        from googleapiclient.errors import HttpError
    except Exception:
        HttpError = Exception  # Fallback for environments with partial installs

    for email in emails:
        if not email:
            continue
        # Template/example addresses should never block the workflow.
        if email.endswith("@example.com"):
            print(f"Skipping placeholder reviewer email: {email}")
            continue
        body = {"type": "user", "role": role, "emailAddress": email}
        try:
            drive.permissions().create(
                fileId=file_id,
                body=body,
                sendNotificationEmail=False,
            ).execute()
        except HttpError as exc:
            # Non-fatal: keep push successful for link-sharing workflows.
            print(f"Warning: could not share with {email}: {exc}")
            continue


def _share_with_anyone_with_link(drive, file_id: str, role: str) -> None:
    """Allow anyone with the link to access the file with a given role."""
    body = {
        "type": "anyone",
        "role": role,
        # False means file is not discoverable by search; link is required.
        "allowFileDiscovery": False,
    }
    drive.permissions().create(fileId=file_id, body=body).execute()


def _list_google_docs_in_folder(drive, folder_id: str) -> list[str]:
    """List Google Doc file IDs inside a folder."""
    query = (
        "mimeType='application/vnd.google-apps.document' and "
        f"trashed=false and '{folder_id}' in parents"
    )
    result = drive.files().list(q=query, fields="files(id,name)", pageSize=200).execute()
    return [f["id"] for f in result.get("files", [])]


def _get_file_name_and_url(drive, file_id: str) -> tuple[str, str]:
    """Return file name and web URL for a Drive file."""
    info = drive.files().get(fileId=file_id, fields="name,webViewLink").execute()
    return info.get("name", ""), info.get("webViewLink", "")


def _fetch_unresolved_comment_summaries(
    drive,
    file_id: str,
    limit: int,
    headings: list[str],
) -> list[dict[str, str]]:
    """Fetch unresolved comment summaries + metadata from a Google Doc."""
    out: list[dict[str, str]] = []
    page_token = None
    seen_pairs: set[tuple[str, str]] = set()
    while True:
        response = drive.comments().list(
            fileId=file_id,
            includeDeleted=False,
            pageSize=100,
            pageToken=page_token,
            fields=(
                "nextPageToken,comments(id,content,anchor,resolved,createdTime,"
                "author(displayName),quotedFileContent,replies(id,content,author(displayName)))"
            ),
        ).execute()
        for c in response.get("comments", []):
            if c.get("resolved", False):
                continue
            content = _normalize_comment_text(c.get("content", "") or "")
            # Avoid recursively carrying forward synthetic carry-forward notes.
            if (
                content.startswith("[FIXED FROM ARCHIVED MANUSCRIPT]")
                or content.startswith("[CARRIED COMMENT FROM PREVIOUS MANUSCRIPT]")
            ):
                continue
            quote = _normalize_comment_text(c.get("quotedFileContent", {}).get("value", ""))
            dedup_key = (content, quote)
            if dedup_key in seen_pairs:
                continue
            seen_pairs.add(dedup_key)
            out.append(
                {
                    "comment_id": c.get("id", ""),
                    "author": c.get("author", {}).get("displayName", "Unknown"),
                    "created_time": c.get("createdTime", ""),
                    "content": content,
                    "quote": quote,
                    "anchor": c.get("anchor", ""),
                    "section_guess": _guess_section(content, quote, headings),
                    "search_key": _derive_search_key(quote),
                }
            )
            if len(out) >= limit:
                return out
        page_token = response.get("nextPageToken")
        if not page_token:
            break
    return out


def _add_fixed_history_comments_to_new_doc(
    drive,
    *,
    new_file_id: str,
    archived_source_name: str,
    archived_source_url: str,
    comment_summaries: list[dict[str, str]],
) -> int:
    """Copy unresolved old comments into the new doc with key metadata."""

    added = 0
    for item in comment_summaries:
        comment_id = (item.get("comment_id", "") or "").strip()
        author = item.get("author", "Unknown")
        created_time = (item.get("created_time", "") or "").strip()
        content = (item.get("content", "") or "(empty)").strip()
        quote = (item.get("quote", "") or "").strip()
        section_guess = (item.get("section_guess", "") or "").strip()
        search_key = (item.get("search_key", "") or "").strip()

        # Keep messages concise and avoid multiline blobs.
        content = re.sub(r"\s+", " ", content)
        quote = re.sub(r"\s+", " ", quote)

        message_lines = [
            "[CARRIED COMMENT FROM PREVIOUS MANUSCRIPT]",
            f"Source: {archived_source_name}",
        ]
        if archived_source_url:
            message_lines.append(f"Archived link: {archived_source_url}")
        if comment_id:
            message_lines.append(f"Original comment ID: {comment_id}")
        message_lines.append(f"Original reviewer: {author}")
        if created_time:
            message_lines.append(f"Original created time: {created_time}")
        message_lines.append(f"Original comment: {content}")
        if section_guess and section_guess != "(Unmapped)":
            message_lines.append(f"Guessed section: {section_guess}")
        if quote:
            message_lines.append(f"Original anchor text: {quote}")
            if search_key:
                message_lines.append(f"Search key phrase: {search_key}")
        message_lines.append("Carry-forward status: unresolved in archived manuscript.")

        drive.comments().create(
            fileId=new_file_id,
            body={"content": "\n".join(message_lines)},
            fields="id",
        ).execute()
        added += 1
    return added


def _move_to_archive(drive, file_id: str, active_folder_id: str, archive_folder_id: str) -> None:
    """Move a file from active review folder to archive folder."""
    drive.files().update(
        fileId=file_id,
        addParents=archive_folder_id,
        removeParents=active_folder_id,
        fields="id,parents",
    ).execute()


def _mark_archived_and_read_only(drive, file_id: str, latest_doc_url: str) -> None:
    """Rename archived docs and downgrade edit/comment access to reader.

    This ensures reviewers immediately see an old doc is archived and cannot
    continue spending effort on outdated content.
    """
    try:
        from googleapiclient.errors import HttpError
    except Exception:
        HttpError = Exception  # Fallback for environments with partial installs

    # 1) Rename + description marker.
    try:
        info = drive.files().get(fileId=file_id, fields="name,description").execute()
        old_name = info.get("name", "")
        if old_name.startswith("[ARCHIVED]"):
            new_name = old_name
        else:
            new_name = f"[ARCHIVED] {old_name}"

        msg = "This manuscript is archived. Please review the latest active manuscript instead."
        if latest_doc_url:
            msg += f" Latest active link: {latest_doc_url}"

        drive.files().update(
            fileId=file_id,
            body={"name": new_name, "description": msg},
            fields="id,name,description",
        ).execute()
    except HttpError as exc:
        print(f"Warning: could not mark archived file metadata for {file_id}: {exc}")

    # 2) Downgrade commenter/writer permissions to reader.
    try:
        perms = drive.permissions().list(
            fileId=file_id,
            fields="permissions(id,type,role,emailAddress,displayName)",
            pageSize=100,
        ).execute().get("permissions", [])
    except HttpError as exc:
        print(f"Warning: could not list permissions for archived file {file_id}: {exc}")
        return

    mutable_types = {"anyone", "user", "group", "domain"}
    for p in perms:
        perm_id = p.get("id")
        p_type = p.get("type", "")
        role = p.get("role", "")
        if not perm_id:
            continue
        if p_type not in mutable_types:
            continue
        if role not in {"writer", "commenter"}:
            continue
        try:
            drive.permissions().update(
                fileId=file_id,
                permissionId=perm_id,
                body={"role": "reader"},
                fields="id,role",
            ).execute()
        except HttpError as exc:
            who = p.get("emailAddress") or p.get("displayName") or p_type
            print(f"Warning: could not downgrade permission for {who} on {file_id}: {exc}")


def _restore_archived_metadata(drive, file_id: str) -> tuple[str, str]:
    """Remove [ARCHIVED] prefix and restore a clean description if present."""
    info = drive.files().get(fileId=file_id, fields="name,description").execute()
    name = info.get("name", "")
    description = info.get("description", "") or ""

    archived_prefix = "[ARCHIVED]"
    new_name = name
    if name.startswith(archived_prefix):
        new_name = name[len(archived_prefix):].lstrip()

    archived_msg = "This manuscript is archived. Please review the latest active manuscript instead."
    new_description = description
    if description.startswith(archived_msg):
        new_description = ""

    if new_name != name or new_description != description:
        drive.files().update(
            fileId=file_id,
            body={"name": new_name, "description": new_description},
            fields="id,name,description",
        ).execute()
    return new_name, new_description


def _restore_permissions(drive, file_id: str, config: ReviewConfig) -> None:
    """Restore reviewer permissions and link-sharing if configured."""
    try:
        from googleapiclient.errors import HttpError
    except Exception:
        HttpError = Exception

    try:
        perms = drive.permissions().list(
            fileId=file_id,
            fields="permissions(id,type,role,emailAddress,displayName)",
            pageSize=100,
        ).execute().get("permissions", [])
    except HttpError as exc:
        print(f"Warning: could not list permissions for {file_id}: {exc}")
        perms = []

    reviewer_emails = {e.lower() for e in config.reviewer_emails if e}
    updated_reviewers: set[str] = set()
    anyone_perm_id = ""
    anyone_role = ""

    for p in perms:
        perm_id = p.get("id")
        p_type = p.get("type", "")
        role = p.get("role", "")
        email = (p.get("emailAddress") or "").lower()

        if p_type == "anyone":
            anyone_perm_id = perm_id or ""
            anyone_role = role
            continue
        if not perm_id or not email:
            continue
        if email not in reviewer_emails:
            continue
        updated_reviewers.add(email)
        if role == config.share_role:
            continue
        try:
            drive.permissions().update(
                fileId=file_id,
                permissionId=perm_id,
                body={"role": config.share_role},
                fields="id,role",
            ).execute()
        except HttpError as exc:
            print(f"Warning: could not restore permission for {email}: {exc}")

    missing_reviewers = [e for e in reviewer_emails if e not in updated_reviewers]
    if missing_reviewers:
        _share_with_reviewers(drive, file_id, missing_reviewers, config.share_role)

    if config.share_with_anyone_with_link:
        if anyone_perm_id and anyone_role != config.anyone_role:
            try:
                drive.permissions().update(
                    fileId=file_id,
                    permissionId=anyone_perm_id,
                    body={"role": config.anyone_role},
                    fields="id,role",
                ).execute()
            except HttpError as exc:
                print(f"Warning: could not restore anyone-link permission: {exc}")
        elif not anyone_perm_id:
            _share_with_anyone_with_link(drive, file_id, config.anyone_role)


def _restore_from_archive(
    drive,
    *,
    file_id: str,
    active_folder_id: str,
    archive_folder_id: str,
    config: ReviewConfig,
) -> tuple[str, str]:
    """Move an archived doc back to ACTIVE and restore permissions."""
    drive.files().update(
        fileId=file_id,
        addParents=active_folder_id,
        removeParents=archive_folder_id,
        fields="id,parents",
    ).execute()

    name, description = _restore_archived_metadata(drive, file_id)
    _restore_permissions(drive, file_id, config)
    return name, description


##### COMMENT FETCH + NORMALIZATION #####

def _extract_headings_from_docx(doc_path: Path) -> list[str]:
    """Extract heading texts from local DOCX for section guessing."""
    try:
        from docx import Document
    except Exception:
        return []

    headings: list[str] = []
    doc = Document(str(doc_path))
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_obj = p.style
        style_name = str(style_obj.name) if style_obj is not None else "Normal"
        if style_name.startswith("Heading"):
            headings.append(text)
    return headings


def _guess_section(comment_text: str, anchor_text: str, headings: list[str]) -> str:
    """Guess section name by heading token overlap."""
    haystack = f"{comment_text} {anchor_text}".lower()
    best = "(Unmapped)"
    best_score = 0
    for heading in headings:
        tokens = [t for t in re.split(r"[^a-z0-9]+", heading.lower()) if len(t) >= 4]
        if not tokens:
            continue
        score = sum(1 for t in tokens if t in haystack)
        if score > best_score:
            best_score = score
            best = heading
    return best


def _extract_docx_section_texts(doc_path: Path) -> list[tuple[str, str]]:
    """Extract (section, paragraph text) pairs from a DOCX for matching."""
    try:
        from docx import Document
    except Exception:
        return []

    out: list[tuple[str, str]] = []
    current_section = "(Unmapped)"
    doc = Document(str(doc_path))
    for paragraph in doc.paragraphs:
        text = _normalize_comment_text(paragraph.text)
        if not text:
            continue
        style_obj = paragraph.style
        style_name = str(style_obj.name) if style_obj is not None else "Normal"
        if style_name.startswith("Heading"):
            current_section = text
            continue
        out.append((current_section, text))
    return out


def _score_text_overlap(left: str, right: str) -> int:
    """Compute overlap score between two text snippets."""
    left_tokens = {t for t in re.split(r"[^a-z0-9]+", left.lower()) if len(t) >= 4}
    right_tokens = {t for t in re.split(r"[^a-z0-9]+", right.lower()) if len(t) >= 4}
    if not left_tokens or not right_tokens:
        return 0
    return len(left_tokens & right_tokens)


def _best_match_section_snippet(
    content: str,
    quote: str,
    search_key: str,
    section_texts: list[tuple[str, str]],
) -> tuple[str, str]:
    """Return best matching section and snippet from the latest manuscript."""
    needle = _normalize_comment_text(" ".join([content, quote, search_key]))
    if not needle or not section_texts:
        return "", ""

    best_section = ""
    best_snippet = ""
    best_score = 0
    for section, snippet in section_texts:
        score = _score_text_overlap(needle, snippet)
        if score > best_score:
            best_score = score
            best_section = section
            best_snippet = snippet
    if best_score == 0:
        return "", ""
    return best_section, best_snippet


def _enrich_comment_summaries_with_doc_match(
    comment_summaries: list[dict[str, str]],
    section_texts: list[tuple[str, str]],
) -> list[dict[str, str]]:
    """Add best-match section/snippet from latest DOCX to comment summaries."""
    enriched: list[dict[str, str]] = []
    for item in comment_summaries:
        section, snippet = _best_match_section_snippet(
            item.get("content", ""),
            item.get("quote", ""),
            item.get("search_key", ""),
            section_texts,
        )
        merged = dict(item)
        if section:
            merged["section_match"] = section
        if snippet:
            merged["snippet_match"] = snippet[:280]
        enriched.append(merged)
    return enriched


def _normalize_comment(
    *,
    version: int,
    index: int,
    raw: dict[str, Any],
    headings: list[str],
) -> ReviewComment:
    """Convert raw API comment thread into dataclass model."""
    quoted = _normalize_comment_text(raw.get("quotedFileContent", {}).get("value", ""))
    content = _normalize_comment_text(raw.get("content", ""))
    section_guess = _guess_section(content, quoted, headings)

    replies: list[CommentReply] = []
    for item in raw.get("replies", []):
        replies.append(
            CommentReply(
                reply_id=item.get("id", ""),
                author=item.get("author", {}).get("displayName", "Unknown"),
                created_time=item.get("createdTime", ""),
                content=_normalize_comment_text(item.get("content", "")),
            )
        )

    return ReviewComment(
        comment_id=raw.get("id", ""),
        stable_id=f"CMT-v{version:03d}-{index:03d}",
        author=raw.get("author", {}).get("displayName", "Unknown"),
        created_time=raw.get("createdTime", ""),
        modified_time=raw.get("modifiedTime", ""),
        resolved=bool(raw.get("resolved", False)),
        content=content,
        quoted_anchor=quoted,
        section_guess=section_guess,
        replies=replies,
    )


def _fetch_comments(drive, file_id: str, version: int, headings: list[str]) -> list[ReviewComment]:
    """Fetch all comment threads and normalize them."""
    out: list[ReviewComment] = []
    page_token = None
    i = 1
    while True:
        response = drive.comments().list(
            fileId=file_id,
            includeDeleted=False,
            pageSize=100,
            pageToken=page_token,
            fields=(
                "nextPageToken,comments(id,content,quotedFileContent,resolved,"
                "createdTime,modifiedTime,author(displayName),"
                "replies(id,content,createdTime,author(displayName)))"
            ),
        ).execute()
        for raw in response.get("comments", []):
            out.append(_normalize_comment(version=version, index=i, raw=raw, headings=headings))
            i += 1
        page_token = response.get("nextPageToken")
        if not page_token:
            break
    return out


def _serialize_comments(comments: list[ReviewComment]) -> list[dict[str, Any]]:
    """Serialize comments dataclasses for JSON output."""
    return [asdict(c) for c in comments]


##### MARKDOWN RENDERING #####

def _render_comments_markdown(
    *,
    version: int,
    comments: list[ReviewComment],
    unresolved_only: bool,
) -> str:
    """Render comments into human-friendly markdown grouped by section."""
    title = "Open" if unresolved_only else "Resolved"
    lines = [
        f"# Reviewer Comments ({title}) - v{version:03d}",
        "",
        f"Generated: {_utc_now()}",
        "",
    ]

    selected = [c for c in comments if (not c.resolved if unresolved_only else c.resolved)]
    if not selected:
        lines.append("No comments in this category.")
        lines.append("")
        return "\n".join(lines) + "\n"

    by_section: dict[str, list[ReviewComment]] = {}
    for item in selected:
        by_section.setdefault(item.section_guess, []).append(item)

    for section in sorted(by_section):
        lines.append(f"## {section}")
        lines.append("")
        for c in by_section[section]:
            checkbox = "- [ ]" if not c.resolved else "- [x]"
            lines.append(f"{checkbox} {c.stable_id} | {c.author} | {c.created_time}")
            lines.append(f"  - Comment: {c.content or '(empty)'}")
            if c.quoted_anchor:
                lines.append(f"  - Anchor: {c.quoted_anchor}")
            if c.replies:
                lines.append("  - Replies:")
                for r in c.replies:
                    lines.append(f"    - {r.author} ({r.created_time}): {r.content}")
            lines.append("")
    return "\n".join(lines) + "\n"


def _render_response_log_entry(ids: list[str], note: str) -> str:
    """Build one response log markdown entry."""
    lines = [f"## {_utc_now()}", ""]
    lines.append("Addressed comment IDs:")
    for cid in ids:
        lines.append(f"- {cid}")
    lines.append("")
    if note:
        lines.append("Notes:")
        lines.append(note)
        lines.append("")
    return "\n".join(lines) + "\n"


##### COMMANDS #####

def _resolve_restore_file_id(
    *,
    version: int | None,
    file_id: str,
    archive_index: int,
) -> tuple[str, Path | None]:
    """Resolve which archived file ID to restore."""
    if file_id:
        return file_id, None
    if version is None:
        raise ValueError("Provide --file-id or --version to restore a doc.")

    review_dir = _review_dir_for_version(version)
    meta_path = review_dir / "gdoc_meta.json"
    if not meta_path.exists():
        raise FileNotFoundError(f"Missing metadata: {meta_path}")

    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    archived_ids = list(meta.get("archived_previous_ids", []))
    if not archived_ids:
        raise ValueError("No archived_previous_ids found in gdoc_meta.json")

    index = max(1, archive_index) - 1
    if index >= len(archived_ids):
        raise ValueError(f"archive_index={archive_index} out of range for archived IDs")

    return archived_ids[index], meta_path


def cmd_restore(args: argparse.Namespace) -> int:
    """Restore an archived Google Doc back to ACTIVE."""
    config = _load_config()
    drive = _get_drive_service(config)

    if not config.drive_parent_folder_id.strip():
        print("drive_parent_folder_id is empty in config.")
        print("Set it in review_bridge_config.json, then retry.")
        return 2

    try:
        file_id, meta_path = _resolve_restore_file_id(
            version=args.version,
            file_id=args.file_id,
            archive_index=args.archive_index,
        )
    except Exception as exc:
        print(f"Restore failed: {exc}")
        return 2

    active_id = _ensure_folder_id(drive, config.active_folder_name, config.drive_parent_folder_id)
    archive_id = _ensure_folder_id(drive, config.archive_folder_name, config.drive_parent_folder_id)

    old_name, old_url = _get_file_name_and_url(drive, file_id)
    new_name, _ = _restore_from_archive(
        drive,
        file_id=file_id,
        active_folder_id=active_id,
        archive_folder_id=archive_id,
        config=config,
    )
    _, new_url = _get_file_name_and_url(drive, file_id)

    restore_meta = {
        "file_id": file_id,
        "restored_at_utc": _utc_now(),
        "old_name": old_name,
        "new_name": new_name,
        "old_url": old_url,
        "new_url": new_url,
    }

    if args.version is not None:
        review_dir = _review_dir_for_version(args.version)
        restore_meta_path = review_dir / "restore_meta.json"
        restore_meta_path.write_text(
            json.dumps(restore_meta, indent=2) + "\n",
            encoding="utf-8",
        )
        print(f"Restore meta: {restore_meta_path}")

    print(f"Restored Google Doc: {new_url}")
    if meta_path is not None:
        print(f"Source meta: {meta_path}")
    return 0

def cmd_push(args: argparse.Namespace) -> int:
    """Upload latest or selected manuscript DOCX as Google Doc for review."""
    config = _load_config()
    version, docx_path = _find_manuscript(args.version)
    review_dir = _review_dir_for_version(version)
    drive = _get_drive_service(config)

    if not config.drive_parent_folder_id.strip():
        print("drive_parent_folder_id is empty in config.")
        print("Set it in review_bridge_config.json, then retry.")
        return 2

    active_id = _ensure_folder_id(drive, config.active_folder_name, config.drive_parent_folder_id)
    archive_id = _ensure_folder_id(drive, config.archive_folder_name, config.drive_parent_folder_id)
    old_active_ids = _list_google_docs_in_folder(drive, active_id)

    new_file_id, new_url = _upload_docx_as_google_doc(drive, docx_path, active_id)
    _share_with_reviewers(drive, new_file_id, config.reviewer_emails, config.share_role)
    if config.share_with_anyone_with_link:
        _share_with_anyone_with_link(drive, new_file_id, config.anyone_role)

    review_copy_path = docx_path.parent / f"{docx_path.stem}_review.docx"
    if review_copy_path.exists():
        review_folder_id = _ensure_folder_id(drive, "REVIEW_COPIES", config.drive_parent_folder_id)
        try:
            review_id, review_url = _upload_docx_as_google_doc(drive, review_copy_path, review_folder_id)
            _share_with_reviewers(drive, review_id, config.reviewer_emails, config.share_role)
            if config.share_with_anyone_with_link:
                _share_with_anyone_with_link(drive, review_id, config.anyone_role)
            review_meta = {
                "review_copy_local": str(review_copy_path),
                "google_file_id": review_id,
                "google_doc_url": review_url,
                "uploaded_at_utc": _utc_now(),
            }
            review_meta_path = review_dir / "review_copy_meta.json"
            review_meta_path.write_text(
                json.dumps(review_meta, indent=2) + "\n",
                encoding="utf-8",
            )
            print(f"Uploaded review copy: {review_url}")
        except Exception as exc:
            print(f"Warning: failed to upload review copy: {exc}")

    carried_forward = 0
    if config.carry_forward_fixed_comments:
        headings = _extract_headings_from_docx(docx_path)
        for fid in old_active_ids:
            if fid == new_file_id:
                continue
            old_name, old_url = _get_file_name_and_url(drive, fid)
            comment_summaries = _fetch_unresolved_comment_summaries(
                drive,
                fid,
                limit=max(1, config.carry_forward_max_comments_per_doc),
                headings=headings,
            )
            if not comment_summaries:
                continue
            added = _add_fixed_history_comments_to_new_doc(
                drive,
                new_file_id=new_file_id,
                archived_source_name=old_name,
                archived_source_url=old_url,
                comment_summaries=comment_summaries,
            )
            carried_forward += added

    archived: list[str] = []
    for fid in old_active_ids:
        if fid == new_file_id:
            continue
        _move_to_archive(drive, fid, active_id, archive_id)
        _mark_archived_and_read_only(drive, fid, new_url)
        archived.append(fid)

    meta = ReviewDocMeta(
        manuscript_version=version,
        local_docx_path=str(docx_path),
        local_version_dir=str(docx_path.parent),
        google_file_id=new_file_id,
        google_doc_url=new_url,
        uploaded_at_utc=_utc_now(),
        archived_previous_ids=archived,
        status="active",
    )
    meta_path = review_dir / "gdoc_meta.json"
    meta_path.write_text(json.dumps(asdict(meta), indent=2) + "\n", encoding="utf-8")

    print(f"Uploaded manuscript v{version:03d}")
    print(f"Google Doc: {new_url}")
    print(f"Meta: {meta_path}")
    print(f"Archived previous docs: {len(archived)}")
    if config.carry_forward_fixed_comments:
        print(f"Carried forward FIXED comments: {carried_forward}")
    return 0


def cmd_download(args: argparse.Namespace) -> int:
    """Download the active Google Doc as DOCX and write gdoc_diff.md."""
    config = _load_config()
    version, docx_path = _find_manuscript(args.version)
    review_dir = _review_dir_for_version(version)
    meta_path = review_dir / "gdoc_meta.json"
    meta = {}
    file_id = ""
    doc_url = ""
    if meta_path.exists():
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        file_id = meta.get("google_file_id", "")
        doc_url = meta.get("google_doc_url", "")

    if not file_id:
        file_id = _extract_drive_id(args.file_id or config.active_file_id)
        if not file_id:
            print(f"Missing metadata: {meta_path}")
            print("Run push first or provide --file-id / active_file_id in config.")
            return 2

    drive = _get_drive_service(config)

    if args.output:
        output_path = Path(args.output)
        if not output_path.is_absolute():
            output_path = project_root / output_path
    else:
        output_path = review_dir / f"gdoc_export_v{version:03d}.docx"

    _export_google_doc_as_docx(drive, file_id, output_path)

    export_meta = {
        "manuscript_version": version,
        "google_file_id": file_id,
        "google_doc_url": doc_url,
        "exported_at_utc": _utc_now(),
        "output_path": str(output_path),
    }
    export_meta_path = review_dir / "gdoc_export_meta.json"
    export_meta_path.write_text(
        json.dumps(export_meta, indent=2) + "\n",
        encoding="utf-8",
    )

    print(f"Downloaded Google Doc to: {output_path}")
    print(f"Export meta: {export_meta_path}")

    if args.skip_diff:
        return 0

    diff_path = review_dir / "gdoc_diff.md"
    diff_text = _build_docx_changelog(docx_path, output_path)
    diff_path.write_text(diff_text + "\n", encoding="utf-8")
    print(f"Wrote Google Doc diff: {diff_path}")

    if config.auto_refresh_review_docx:
        try:
            review_path = docx_path.parent / f"{docx_path.stem}_review.docx"
            _generate_review_docx(output_path, docx_path, review_path)
            print(f"Refreshed review DOCX: {review_path}")
        except Exception as exc:
            print(f"Warning: failed to refresh review DOCX: {exc}")
    return 0


def cmd_review_docx(args: argparse.Namespace) -> int:
    """Generate review DOCX by diffing ACTIVE Google Doc vs local manuscript."""
    version, docx_path = _find_manuscript(args.version)
    try:
        if args.file_id:
            config = _load_config()
            review_dir = _review_dir_for_version(version)
            export_path = review_dir / f"gdoc_export_v{version:03d}.docx"
            drive = _get_drive_service(config)
            file_id = _extract_drive_id(args.file_id)
            _export_google_doc_as_docx(drive, file_id, export_path)
            review_path = docx_path.parent / f"{docx_path.stem}_review.docx"
            _generate_review_docx(export_path, docx_path, review_path)
        else:
            review_path = _build_review_docx_from_active(version, docx_path)
    except Exception as exc:
        print(f"Review DOCX generation failed: {exc}")
        return 2
    print(f"Generated review DOCX: {review_path}")
    return 0


def cmd_pull(args: argparse.Namespace) -> int:
    """Fetch comments from Google Docs and write local review files."""
    config = _load_config()
    version, docx_path = _find_manuscript(args.version)
    review_dir = _review_dir_for_version(version)
    meta_path = review_dir / "gdoc_meta.json"
    if not meta_path.exists():
        print(f"Missing metadata: {meta_path}")
        print("Run push first.")
        return 2

    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    file_id = meta.get("google_file_id", "")
    if not file_id:
        print("google_file_id missing in metadata.")
        return 2

    drive = _get_drive_service(config)
    headings = _extract_headings_from_docx(docx_path)
    comments = _fetch_comments(drive, file_id, version, headings)

    raw_path = review_dir / "comments_raw.json"
    open_md = review_dir / "comments_open.md"
    done_md = review_dir / "comments_resolved.md"

    raw_path.write_text(
        json.dumps(_serialize_comments(comments), indent=2) + "\n",
        encoding="utf-8",
    )
    open_md.write_text(
        _render_comments_markdown(version=version, comments=comments, unresolved_only=True),
        encoding="utf-8",
    )
    done_md.write_text(
        _render_comments_markdown(version=version, comments=comments, unresolved_only=False),
        encoding="utf-8",
    )

    unresolved = sum(1 for c in comments if not c.resolved)
    resolved = sum(1 for c in comments if c.resolved)
    print(f"Pulled comments for manuscript v{version:03d}")
    print(f"Unresolved: {unresolved} | Resolved: {resolved}")
    print(f"Raw: {raw_path}")
    print(f"Open: {open_md}")
    print(f"Resolved: {done_md}")
    return 0


def _build_local_changelog(version: int) -> str:
    """Build markdown changelog between previous and current manuscript."""
    try:
        import manuscript_changelog as mch
    except Exception as exc:
        return f"Could not load manuscript_changelog.py: {exc}\n"

    all_versions = sorted(_collect_existing_manuscripts())
    if version not in all_versions:
        return "Current manuscript version not found.\n"
    idx = all_versions.index(version)
    if idx == 0:
        return "No previous manuscript exists for changelog.\n"

    old_v = all_versions[idx - 1]
    old_path = _collect_existing_manuscripts()[old_v]
    new_path = _collect_existing_manuscripts()[version]

    old_paras = mch.extract_paragraphs(old_path)
    new_paras = mch.extract_paragraphs(new_path)
    changelog = mch.compute_changelog(old_paras, new_paras, old_path, new_path)
    return mch.format_changelog(changelog)


def _build_docx_changelog(old_path: Path, new_path: Path) -> str:
    """Build markdown changelog between two arbitrary DOCX files."""
    try:
        import manuscript_changelog as mch
    except Exception as exc:
        return f"Could not load manuscript_changelog.py: {exc}\n"

    old_paras = mch.extract_paragraphs(old_path)
    new_paras = mch.extract_paragraphs(new_path)
    changelog = mch.compute_changelog(old_paras, new_paras, old_path, new_path)
    return mch.format_changelog(changelog)


##### REVIEW DOCX DIFF ######

def _collect_doc_paragraphs(doc: Document) -> list[tuple[Paragraph, str, str]]:
    """Collect non-empty paragraphs with their styles."""
    items: list[tuple[Paragraph, str, str]] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_obj = p.style
        style_name = str(style_obj.name) if style_obj is not None else "Normal"
        items.append((p, text, style_name))
    return items


def _clear_paragraph_runs(paragraph: Paragraph) -> None:
    """Remove all runs from a paragraph."""
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def _apply_review_markup(paragraph: Paragraph, text: str) -> None:
    """Apply **++insert++**, ~~delete~~, and 💬 [COMMENT] markup to a paragraph."""
    _clear_paragraph_runs(paragraph)
    tokens = re.split(r'(\*\*\+\+.*?\+\+\*\*|~~.*?~~|💬 \[.*?\])', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**++") and token.endswith("++**"):
            run = paragraph.add_run(token[4:-4])
            run.font.color.rgb = RGBColor(0, 128, 0)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif token.startswith("~~") and token.endswith("~~"):
            run = paragraph.add_run(token[2:-2])
            run.font.color.rgb = RGBColor(255, 0, 0)
            run.font.strike = True
        elif token.startswith("💬 ["):
            run = paragraph.add_run(" " + token + " ")
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
        else:
            paragraph.add_run(token)


def _word_diff_review(old: str, new: str) -> str:
    """Return word-level diff with **++ ++** for insertions and ~~ ~~ for deletes."""
    old_words = old.split()
    new_words = new.split()
    matcher = difflib.SequenceMatcher(None, old_words, new_words)
    parts: list[str] = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            parts.append(" ".join(old_words[i1:i2]))
        elif tag == "replace":
            parts.append("~~" + " ".join(old_words[i1:i2]) + "~~")
            parts.append("**++" + " ".join(new_words[j1:j2]) + "++**")
        elif tag == "insert":
            parts.append("**++" + " ".join(new_words[j1:j2]) + "++**")
        elif tag == "delete":
            parts.append("~~" + " ".join(old_words[i1:i2]) + "~~")
    return " ".join(p for p in parts if p)


def _insert_paragraph_after(anchor: Paragraph, style: str | None = None) -> Paragraph:
    """Insert a paragraph after anchor and return it."""
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    new_para = Paragraph(new_p, anchor._parent)
    if style:
        new_para.style = style
    return new_para


def _insert_paragraph_before(anchor: Paragraph, style: str | None = None) -> Paragraph:
    """Insert a paragraph before anchor and return it."""
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    new_para = Paragraph(new_p, anchor._parent)
    if style:
        new_para.style = style
    return new_para


def _style_name(paragraph: Paragraph | None) -> str | None:
    """Return style name for paragraph if available."""
    if paragraph is None:
        return None
    style_obj = paragraph.style
    return str(style_obj.name) if style_obj is not None else None


def _insert_author_note_top(doc: Document, text: str) -> None:
    """Insert a highlighted author note at the top of the document."""
    if not doc.paragraphs:
        return
    p = _insert_paragraph_before(doc.paragraphs[0])
    run = p.add_run(f"[Author Note: {text}]")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
    shading_elem = OxmlElement("w:shd")
    shading_elem.set(qn("w:val"), "clear")
    shading_elem.set(qn("w:color"), "auto")
    shading_elem.set(qn("w:fill"), "E6F2FF")
    p_pr = p._p.get_or_add_pPr()
    p_pr.append(shading_elem)
    ind_elem = OxmlElement("w:ind")
    ind_elem.set(qn("w:left"), "360")
    ind_elem.set(qn("w:right"), "360")
    p_pr.append(ind_elem)


def _generate_review_docx(old_path: Path, new_path: Path, output_path: Path) -> None:
    """Create a review-marked DOCX by diffing old vs new manuscripts."""
    old_doc = Document(str(old_path))
    new_doc = Document(str(new_path))

    old_items = [p.text.strip() for p in old_doc.paragraphs if p.text.strip()]
    new_items = _collect_doc_paragraphs(new_doc)
    new_texts = [t for _, t, _ in new_items]

    matcher = difflib.SequenceMatcher(None, old_items, new_texts)
    last_anchor: Paragraph | None = None

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for idx in range(j1, j2):
                last_anchor = new_items[idx][0]
        elif tag == "insert":
            for idx in range(j1, j2):
                para, text, _ = new_items[idx]
                _apply_review_markup(para, f"**++{text}++**")
                last_anchor = para
        elif tag == "delete":
            for old_text in old_items[i1:i2]:
                style = _style_name(last_anchor)
                if last_anchor is None and new_items:
                    last_anchor = _insert_paragraph_before(new_items[0][0], style=style)
                else:
                    last_anchor = _insert_paragraph_after(last_anchor, style=style)
                _apply_review_markup(last_anchor, f"~~{old_text}~~")
        elif tag == "replace":
            old_slice = old_items[i1:i2]
            new_slice = new_items[j1:j2]
            max_len = max(len(old_slice), len(new_slice))
            for k in range(max_len):
                old_text = old_slice[k] if k < len(old_slice) else None
                if k < len(new_slice):
                    para, new_text, _ = new_slice[k]
                else:
                    para, new_text = None, None
                if old_text and new_text and para is not None:
                    diff_text = _word_diff_review(old_text, new_text)
                    _apply_review_markup(para, diff_text)
                    last_anchor = para
                elif new_text and para is not None:
                    _apply_review_markup(para, f"**++{new_text}++**")
                    last_anchor = para
                elif old_text:
                    style = _style_name(last_anchor)
                    if last_anchor is None and new_items:
                        last_anchor = _insert_paragraph_before(new_items[0][0], style=style)
                    else:
                        last_anchor = _insert_paragraph_after(last_anchor, style=style)
                    _apply_review_markup(last_anchor, f"~~{old_text}~~")

    _insert_author_note_top(
        new_doc,
        "hocam bu review dosyasi aktif gdoc ile son manuskrip arasindaki farklari otomatik gosterir."
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    new_doc.save(str(output_path))


def _build_review_docx_from_active(version: int, local_docx_path: Path) -> Path:
    """Download ACTIVE Google Doc and generate review-marked DOCX."""
    config = _load_config()
    review_dir = _review_dir_for_version(version)
    meta_path = review_dir / "gdoc_meta.json"
    file_id = ""
    if meta_path.exists():
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        file_id = meta.get("google_file_id", "")

    if not file_id:
        file_id = _extract_drive_id(config.active_file_id)
    if not file_id:
        raise RuntimeError(f"Missing metadata: {meta_path}")

    drive = _get_drive_service(config)
    export_path = review_dir / f"gdoc_export_v{version:03d}.docx"
    _export_google_doc_as_docx(drive, file_id, export_path)

    output_path = local_docx_path.parent / f"{local_docx_path.stem}_review.docx"
    _generate_review_docx(export_path, local_docx_path, output_path)
    return output_path


def _build_gdoc_diff_from_active(
    version: int,
    local_docx_path: Path,
    output_path: Path | None = None,
) -> tuple[Path, bool]:
    """Download ACTIVE Google Doc and write a changelog diff against local."""
    config = _load_config()
    review_dir = _review_dir_for_version(version)
    meta_path = review_dir / "gdoc_meta.json"
    file_id = ""
    if meta_path.exists():
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        file_id = meta.get("google_file_id", "")

    if not file_id:
        file_id = _extract_drive_id(config.active_file_id)
    if not file_id:
        raise RuntimeError(f"Missing metadata: {meta_path}")

    drive = _get_drive_service(config)
    export_path = review_dir / f"gdoc_export_v{version:03d}.docx"
    _export_google_doc_as_docx(drive, file_id, export_path)

    diff_text = _build_docx_changelog(local_docx_path, export_path)
    diff_path = output_path or (review_dir / "gdoc_diff.md")
    diff_path.write_text(diff_text + "\n", encoding="utf-8")

    has_changes = "No changes detected" not in diff_text
    return diff_path, has_changes


def cmd_packet(args: argparse.Namespace) -> int:
    """Build a single action packet from comments + local manuscript diff."""
    version, _ = _find_manuscript(args.version)
    review_dir = _review_dir_for_version(version)
    open_md = review_dir / "comments_open.md"
    packet_path = review_dir / "review_packet.md"

    if not open_md.exists():
        print(f"Missing open comments file: {open_md}")
        print("Run pull first.")
        return 2

    comments_text = _safe_read_text(open_md)
    changelog_text = _build_local_changelog(version)
    gdoc_diff_path = review_dir / "gdoc_diff.md"
    gdoc_diff_text = _safe_read_text(gdoc_diff_path)

    gdoc_section = gdoc_diff_text.strip() if gdoc_diff_text else (
        "_No Google Docs diff found. Run `download` to generate gdoc_diff.md._"
    )

    lines = [
        f"# Review Packet - v{version:03d}",
        "",
        f"Generated: {_utc_now()}",
        "",
        "## 1) Open Reviewer Comments",
        "",
        comments_text.strip(),
        "",
        "## 2) Google Docs vs Local Manuscript Diff",
        "",
        gdoc_section,
        "",
        "## 3) Local Manuscript Changelog",
        "",
        changelog_text.strip(),
        "",
        "## 4) Resolution Checklist",
        "",
        "- [ ] Convert each open comment into one concrete manuscript edit",
        "- [ ] Rebuild manuscript DOCX",
        "- [ ] Re-run pull after reviewers add new feedback",
        "- [ ] Append addressed IDs via close command",
        "",
    ]
    packet_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(f"Built review packet: {packet_path}")
    return 0


def cmd_close(args: argparse.Namespace) -> int:
    """Append an addressed-comments note to response log."""
    version, _ = _find_manuscript(args.version)
    review_dir = _review_dir_for_version(version)
    log_path = review_dir / "response_log.md"

    ids: list[str] = list(args.ids or [])
    if not ids and args.ids_file:
        ids = [line.strip() for line in _safe_read_text(Path(args.ids_file)).splitlines() if line.strip()]
    if not ids:
        print("No comment IDs provided. Use --ids or --ids-file.")
        return 2

    if not log_path.exists():
        log_path.write_text(f"# Response Log - v{version:03d}\n\n", encoding="utf-8")

    entry = _render_response_log_entry(ids, args.note or "")
    with log_path.open("a", encoding="utf-8") as f:
        f.write(entry)

    print(f"Appended response log entry: {log_path}")
    print(f"Recorded IDs: {', '.join(ids)}")
    return 0


##### CLI #####

def _build_parser() -> argparse.ArgumentParser:
    """Create command-line argument parser."""
    parser = argparse.ArgumentParser(
        description="Automate Google Docs review loop for versioned GSM manuscripts"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("init-config", help="Create starter review_bridge_config.json")

    p_push = sub.add_parser("push", help="Upload manuscript to Google Docs and archive old")
    p_push.add_argument("--version", type=int, default=None,
                        help="Manuscript version number (default: latest)")

    p_restore = sub.add_parser("restore", help="Restore an archived Google Doc to ACTIVE")
    p_restore.add_argument("--file-id", type=str, default="",
                           help="Archived Google Doc file ID to restore")
    p_restore.add_argument("--version", type=int, default=None,
                           help="Manuscript version for gdoc_meta.json lookup")
    p_restore.add_argument("--archive-index", type=int, default=1,
                           help="Index into archived_previous_ids (1-based)")

    p_download = sub.add_parser("download", help="Export Google Doc to DOCX and diff")
    p_download.add_argument("--version", type=int, default=None,
                            help="Manuscript version number (default: latest)")
    p_download.add_argument("--output", type=str, default="",
                            help="Optional output DOCX path (default: review folder)")
    p_download.add_argument("--skip-diff", action="store_true",
                            help="Skip gdoc_diff.md generation")
    p_download.add_argument("--file-id", type=str, default="",
                            help="Optional Google Doc file id or URL")

    p_review = sub.add_parser("review-docx", help="Generate review DOCX from ACTIVE Google Doc")
    p_review.add_argument("--version", type=int, default=None,
                          help="Manuscript version number (default: latest)")
    p_review.add_argument("--file-id", type=str, default="",
                          help="Optional Google Doc file id or URL")

    p_pull = sub.add_parser("pull", help="Fetch comments for manuscript version")
    p_pull.add_argument("--version", type=int, default=None,
                        help="Manuscript version number (default: latest)")

    p_packet = sub.add_parser("packet", help="Generate combined review packet markdown")
    p_packet.add_argument("--version", type=int, default=None,
                          help="Manuscript version number (default: latest)")

    p_close = sub.add_parser("close", help="Log addressed comment IDs")
    p_close.add_argument("--version", type=int, default=None,
                         help="Manuscript version number (default: latest)")
    p_close.add_argument("--ids", nargs="*", default=[],
                         help="Comment IDs like CMT-v034-001")
    p_close.add_argument("--ids-file", type=str, default="",
                         help="Path to text file containing one comment ID per line")
    p_close.add_argument("--note", type=str, default="",
                         help="Optional free-text note for this closure event")

    return parser


def main() -> int:
    """CLI entry point."""
    parser = _build_parser()
    args = parser.parse_args()

    if args.command == "init-config":
        return cmd_init_config(args)
    if args.command == "push":
        return cmd_push(args)
    if args.command == "restore":
        return cmd_restore(args)
    if args.command == "download":
        return cmd_download(args)
    if args.command == "pull":
        return cmd_pull(args)
    if args.command == "packet":
        return cmd_packet(args)
    if args.command == "close":
        return cmd_close(args)
    if args.command == "review-docx":
        return cmd_review_docx(args)

    parser.print_help()
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
